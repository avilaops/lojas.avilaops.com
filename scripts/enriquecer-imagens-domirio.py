from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


ORIGEM = Path("output/imports/domirio-produtos-merchant-v7.xlsx")
DESTINO = Path("output/imports/domirio-produtos-imagens-v8.xlsx")

IMAGENS_OFICIAIS = {
    "Cypertryne": "https://isorgan.com.br/images/produtos/galeria/cypertryne/1.webp",
    "Glifosato 480": "https://isorgan.com.br/images/produtos/galeria/glifosato-480/1.webp",
    "Hora 2": "https://isorgan.com.br/images/produtos/galeria/hora-2/1.webp",
    "Pirazol TOP": "https://isorgan.com.br/images/produtos/galeria/pirazol-top/1.webp",
    "Roetox": "https://isorgan.com.br/images/produtos/galeria/roetox/1.webp",
    "Sanepec": "https://isorgan.com.br/images/produtos/galeria/sanepec/1.webp",
    "Scorpfim": "https://isorgan.com.br/images/produtos/galeria/scorpfim/1.webp",
    "Termitox 40": "https://isorgan.com.br/images/produtos/galeria/termitox-40/1.webp",
    "Termitox 400": "https://isorgan.com.br/images/produtos/galeria/termitox-400/1.webp",
    "Termitox Spray": "https://isorgan.com.br/images/produtos/galeria/termitox-spray/1.webp",
    "Termitox Spray DUO": "https://isorgan.com.br/images/produtos/galeria/termitox-spray-duo/1.webp",
}


def main() -> None:
    workbook = load_workbook(ORIGEM)
    sheet = workbook.active
    headers = {cell.value: cell.column for cell in sheet[1]}
    preenchidos = 0

    for row in range(2, sheet.max_row + 1):
        if str(sheet.cell(row, headers["ativo"]).value).strip().lower() != "sim":
            continue
        familia = str(sheet.cell(row, headers["imagem_familia"]).value or "").strip()
        if familia not in IMAGENS_OFICIAIS:
            import json

            atributos = json.loads(sheet.cell(row, headers["atributos_json"]).value or "{}")
            familia = str(atributos.get("familia") or "").strip()
        imagem = IMAGENS_OFICIAIS.get(familia)
        if not imagem:
            continue

        sheet.cell(row, headers["imagem"], imagem)
        sheet.cell(row, headers["imagens"], imagem)
        sheet.cell(row, headers["imagem_origem"], "representativa")
        sheet.cell(row, headers["imagem_familia"], familia)
        sheet.cell(row, headers["confirmar_imagem_exata"], "nao")
        sheet.cell(row, headers["correspondencia_imagem"], "nao_confirmada")
        preenchidos += 1

    workbook.save(DESTINO)
    print(f"arquivo={DESTINO.resolve()}")
    print(f"imagens_oficiais={preenchidos}")


if __name__ == "__main__":
    main()
